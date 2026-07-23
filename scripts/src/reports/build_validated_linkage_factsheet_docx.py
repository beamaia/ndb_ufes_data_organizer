#!/usr/bin/env python3
"""Build editable DOCX factsheets for the validated patch-to-WSI linkage."""

import argparse
import csv
import hashlib
import json
from collections import Counter, defaultdict
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from PIL import Image, ImageDraw


ROOT = Path(__file__).resolve().parents[1]
VALIDATED_DIR = ROOT / "results/phase0/validated_linkage"
PRIVATE_DIR = VALIDATED_DIR / "private_lab_crosswalks"
PATCH_DIR = ROOT / "data/ndb_ufes/patch_level/images"
FIGURE_DIR = VALIDATED_DIR / "figures"
DOCX_CACHE = VALIDATED_DIR / "docx_render_cache"
PDF_IMAGE_CACHE = VALIDATED_DIR / "render_cache"

LINKAGE_CSV = VALIDATED_DIR / "validated_patch_wsi_linkage.csv"
PRIVATE_LINKAGE_CSV = PRIVATE_DIR / "private_validated_patch_wsi_linkage.csv"
INVENTORY_CSV = VALIDATED_DIR / "validated_wsi_inventory.csv"
SIMILARITY_CSV = VALIDATED_DIR / "validated_wsi_patch_similarity.csv"
CROSSWALK_CSV = PRIVATE_DIR / "private_validated_wsi_crosswalk.csv"
SUMMARY_JSON = VALIDATED_DIR / "validated_linkage_summary.json"

LABEL_DISPLAY = {
    "oscc": "OSCC",
    "with_dysplasia": "with dysplasia",
    "without_dysplasia": "without dysplasia",
    "broad_leukoplakia_unknown_dysplasia": "broad leukoplakia, dysplasia unknown",
    "unknown": "not available",
    "": "not available",
}

LABEL_COLORS = {
    "oscc": (208, 74, 58),
    "with_dysplasia": (226, 162, 58),
    "without_dysplasia": (62, 120, 178),
    "unknown": (110, 110, 110),
    "": (110, 110, 110),
}

EVIDENCE_DISPLAY = {
    "validated_exact_public_wsi": "Exact patch-to-WSI link and public WSI match",
    "validated_exact_sab_only": "Exact patch-to-WSI link, SAB-only WSI",
    "validated_exact_with_metadata_conflict": "Exact patch-to-WSI link with WSI context review flag",
    "requires_manual_review": "Requires manual review",
}

PLAUSIBILITY_DISPLAY = {
    "plausible": "Patch label is plausible in the validated WSI context",
    "plausible_but_origin_patch_disagreement": "Patch label differs from WSI folder label but remains biologically plausible",
    "impossible_or_high_conflict": "Patch label is not plausible for the WSI folder label and requires expert review",
    "broad_origin_requires_review": "Broad leukoplakia WSI context requires review",
    "unknown_origin_metadata": "WSI context is not available",
}

WSI_SOURCE_DISPLAY = {
    "both": "public NDB-UFES + SAB",
    "sab_only_recovered": "SAB-only recovered WSI",
    "SAB-only recovered WSI": "SAB-only recovered WSI",
}

CONTENT_WIDTH_CM = 17.0
TABLE_WIDTH_DXA = 9638


class Numbering:
    def __init__(self):
        self.table = 0
        self.figure = 0

    def table_caption(self, title):
        self.table += 1
        return f"Table {self.table}. {title}"

    def figure_caption(self, title):
        self.figure += 1
        return f"Figure {self.figure}. {title}"


def read_csv_rows(path):
    with Path(path).open(newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def read_json(path):
    return json.loads(Path(path).read_text(encoding="utf-8"))


def display_label(value):
    return LABEL_DISPLAY.get(str(value or ""), str(value or "not available").replace("_", " "))


def display_source(value):
    return WSI_SOURCE_DISPLAY.get(str(value or ""), str(value or "not available").replace("_", " "))


def safe_float(value, default=0.0):
    try:
        if value in (None, "", "nan"):
            return default
        return float(value)
    except (TypeError, ValueError):
        return default


def safe_int_text(value):
    if value in (None, "", "nan"):
        return "not available"
    try:
        return str(int(float(value)))
    except (TypeError, ValueError):
        return str(value)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_width(table, widths_dxa):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            if idx < len(widths_dxa):
                set_cell_width(cell, widths_dxa[idx])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_paragraph_keep_with_next(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    keep_next = p_pr.find(qn("w:keepNext"))
    if keep_next is None:
        keep_next = OxmlElement("w:keepNext")
        p_pr.append(keep_next)


def style_cell_text(cell, bold=False, size=8.5, color=None):
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_after = Pt(0)
        for run in paragraph.runs:
            run.font.name = "Arial"
            run.font.size = Pt(size)
            run.bold = bold
            if color:
                run.font.color.rgb = RGBColor(*color)


def add_table(doc, rows, widths, header=True):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    set_table_width(table, widths)
    for r_idx, row_values in enumerate(rows):
        for c_idx, value in enumerate(row_values):
            cell = table.cell(r_idx, c_idx)
            cell.text = str(value)
            style_cell_text(cell, bold=header and r_idx == 0, size=8.5)
            if header and r_idx == 0:
                set_cell_shading(cell, "F2F4F7")
        if header and r_idx == 0:
            set_repeat_table_header(table.rows[0])
    doc.add_paragraph()
    return table


def add_caption(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.style = "Caption"
    run = paragraph.add_run(text)
    run.bold = True
    return paragraph


def add_body(doc, text):
    p = doc.add_paragraph(text)
    p.style = "Body Text"
    return p


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    set_paragraph_keep_with_next(p)
    return p


def add_static_toc(doc, entries):
    add_heading(doc, "Table of contents", 1)
    add_body(doc, "This editable DOCX uses Word heading styles. In Microsoft Word or Word Online, use the references/tools menu to insert or update a generated table of contents if needed.")
    for title, level in entries:
        p = doc.add_paragraph()
        p.style = "Body Text"
        if level == 2:
            p.paragraph_format.left_indent = Cm(0.6)
        p.add_run(title)


def count_rows(rows, column, display=None):
    counter = Counter((row.get(column) or "unknown") for row in rows)
    total = sum(counter.values()) or 1
    out = [["Category", "Count", "Percent"]]
    for key, count in counter.most_common():
        label = display.get(key, key.replace("_", " ")) if display else key.replace("_", " ")
        out.append([label, f"{count:,}", f"{count / total * 100:.1f}%"])
    return out


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)

    for style_name, size, color, before, after in [
        ("Title", 24, (0, 0, 0), 0, 6),
        ("Heading 1", 18, (0, 0, 0), 18, 7),
        ("Heading 2", 14, (0, 0, 0), 14, 5),
        ("Heading 3", 12, (67, 67, 67), 10, 4),
        ("Body Text", 10.5, (0, 0, 0), 0, 6),
        ("Caption", 8.5, (85, 85, 85), 3, 7),
    ]:
        style = styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(*color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.12


def add_cover(doc, summary, privacy):
    title = doc.add_paragraph()
    title.style = "Title"
    title.add_run("Validated dataset factsheet").bold = False
    subtitle = doc.add_paragraph()
    subtitle.style = "Title"
    subtitle.add_run("NDB-UFES and SAB patch-to-WSI linkage")
    add_body(doc, "Editable DOCX source document")
    add_body(doc, f"{int(summary['patches']):,} patches | {int(summary['validated_wsi_count']):,} validated WSI units | {privacy} version")
    add_body(doc, "Pixel containment, label agreement, WSI inventory, patch similarity, and full WSI atlas.")
    doc.add_page_break()


def add_intro_sections(doc, numbering, summary, linkage_rows, inventory_rows):
    add_heading(doc, "Validated linkage summary", 1)
    add_body(doc, "This document describes the repaired patch-to-WSI linkage between the public NDB-UFES dataset and the private SAB laboratory dataset. The WSI assignment is based on exact pixel containment of each patch inside the SAB WSI image.")
    overview = [
        ["Measure", "Value"],
        ["Patch images linked to validated WSI IDs", f"{int(summary['patches_with_final_wsi_id']):,} / {int(summary['patches']):,}"],
        ["Patch images with recovered coordinates", f"{int(summary['patches_with_coordinates']):,} / {int(summary['patches']):,}"],
        ["Validated WSI units", f"{int(summary['validated_wsi_count']):,}"],
        ["Validated WSI units with public NDB-UFES exact WSI match", f"{int(summary['validated_wsi_with_public_ndb_match']):,}"],
        ["SAB-only recovered WSI units", f"{int(summary['validated_wsi_sab_only']):,}"],
    ]
    add_caption(doc, numbering.table_caption("Validated linkage overview"))
    add_table(doc, overview, [6200, 3200])

    add_heading(doc, "Dataset sources", 1)
    add_body(doc, "NDB-UFES is the public Mendeley dataset. SAB is the private laboratory dataset used to recover patch provenance and exact WSI containment.")
    sources = [
        ["Source", "Description", "Role in this document"],
        ["NDB-UFES", "Public dataset distributed through Mendeley.", "Provides public WSI images, patch metadata, and public dataset context."],
        ["SAB", "Private laboratory dataset with WSI class folders and patch train/test class folders.", "Provides the WSI files and patch folders used to recover exact patch containment."],
        ["Validated WSI ID", "A public pseudonym assigned to each final WSI unit.", "Canonical WSI identity used in the editable document."],
        ["Private crosswalk", "Laboratory-only table linking validated WSI IDs to SAB names and paths.", "Included only in the lab version."],
    ]
    add_caption(doc, numbering.table_caption("Dataset sources and roles"))
    add_table(doc, sources, [2100, 3600, 3700])

    add_heading(doc, "Linkage evidence and patch labels", 1)
    add_body(doc, "NDB-UFES and SAB patch labels agree for all 3,763 patches. The main product of this document is therefore the validated patch-to-WSI linkage, not a patch-label correction.")
    add_caption(doc, numbering.table_caption("Linkage evidence levels"))
    add_table(doc, count_rows(linkage_rows, "linkage_evidence_level", EVIDENCE_DISPLAY), [6200, 1600, 1600])
    add_caption(doc, numbering.table_caption("NDB-UFES patch label versus SAB patch label"))
    add_table(doc, count_rows(linkage_rows, "patch_label_agreement_status", {"agrees": "NDB-UFES patch label and SAB patch label agree", "disagrees": "NDB-UFES patch label and SAB patch label differ"}), [6200, 1600, 1600])
    add_caption(doc, numbering.table_caption("Patch label versus validated WSI context"))
    add_table(doc, count_rows(linkage_rows, "patch_vs_origin_plausibility_status", PLAUSIBILITY_DISPLAY), [6200, 1600, 1600])

    add_heading(doc, "Validated WSI inventory", 1)
    source_counts = Counter(row.get("wsi_source", "") for row in inventory_rows)
    source_table = [["WSI source", "Number of WSI units"]]
    for key, count in source_counts.most_common():
        source_table.append([display_source(key), str(count)])
    add_caption(doc, numbering.table_caption("Validated WSI source summary"))
    add_table(doc, source_table, [6200, 3200])

    bins = [
        ("1 patch", lambda n: n == 1),
        ("2 to 5 patches", lambda n: 2 <= n <= 5),
        ("6 to 10 patches", lambda n: 6 <= n <= 10),
        ("11 to 15 patches", lambda n: 11 <= n <= 15),
        ("16 to 20 patches", lambda n: 16 <= n <= 20),
    ]
    patch_count_rows = [["Patch count group", "Number of WSI units"]]
    for label, predicate in bins:
        patch_count_rows.append([label, str(sum(1 for row in inventory_rows if predicate(int(row.get("n_patches") or 0))))])
    add_caption(doc, numbering.table_caption("Number of patches per validated WSI"))
    add_table(doc, patch_count_rows, [6200, 3200])


def add_similarity_sections(doc, numbering):
    add_heading(doc, "Same-WSI patch similarity", 1)
    add_body(doc, "Same-WSI patch similarity is descriptive. It helps identify whether patches from the same validated WSI overlap spatially, share similar visual fingerprints, or provide different tissue information.")
    methods = [
        ["Measure", "Definition", "How to read it"],
        ["IoU", "Intersection over union between two recovered coordinate boxes in the same validated WSI.", "Higher values mean more spatial overlap."],
        ["Visual fingerprint similarity", "Patch images are resized to 32 x 32 pixels, RGB and LAB vectors are concatenated, mean-centered, L2-normalized, and compared by cosine similarity.", "Higher values mean more similar color and texture fingerprints. It is not a diagnosis score."],
        ["LAB mean delta", "Euclidean distance between mean LAB color values of the two patches.", "Lower values mean more similar average stain/color appearance."],
        ["Pair relation", "Readable category derived from spatial overlap, visual fingerprint similarity, and LAB color difference.", "Separates overlapping patches from visually similar but spatially separate patches."],
    ]
    add_caption(doc, numbering.table_caption("Same-WSI patch similarity measures"))
    add_table(doc, methods, [2100, 4700, 2600])

    figure_paths = [
        ("same_wsi_iou_distribution.png", "Same-WSI spatial overlap distribution"),
        ("same_wsi_visual_similarity_distribution.png", "Same-WSI visual fingerprint similarity distribution"),
        ("same_wsi_lab_delta_distribution.png", "Same-WSI LAB color difference distribution"),
        ("same_wsi_pair_relation_counts.png", "Same-WSI pair interpretation counts"),
    ]
    for file_name, title in figure_paths:
        path = FIGURE_DIR / file_name
        if path.exists():
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(path), width=Cm(14.0))
            add_caption(doc, numbering.figure_caption(title))


def make_context_image(wsi_id, rows):
    DOCX_CACHE.mkdir(parents=True, exist_ok=True)
    out = DOCX_CACHE / f"{wsi_id}_context.jpg"
    source = rows[0].get("sab_origin_path") or ""
    if not source or not Path(source).exists():
        return None
    source_path = Path(source)
    if out.exists():
        return out
    with Image.open(source_path) as original:
        original_w, original_h = original.size
    pdf_cache_key = hashlib.sha1(f"{source_path.resolve()}|1100|82".encode("utf-8")).hexdigest()[:24]
    cached_preview = PDF_IMAGE_CACHE / f"{pdf_cache_key}.jpg"
    if cached_preview.exists():
        with Image.open(cached_preview) as img:
            img = img.convert("RGB")
            display_w, display_h = img.size
    else:
        with Image.open(source_path) as img:
            img = img.convert("RGB")
            img.thumbnail((1100, 1100), Image.Resampling.LANCZOS)
            display_w, display_h = img.size
    scale = min(display_w / original_w, display_h / original_h)
    draw = ImageDraw.Draw(img)
    for row in rows:
        label = row.get("sab_patch_label") or row.get("ndb_ufes_patch_label") or "unknown"
        color = LABEL_COLORS.get(label, LABEL_COLORS["unknown"])
        x = safe_float(row.get("x")) * scale
        y = safe_float(row.get("y")) * scale
        w = safe_float(row.get("width")) * scale
        h = safe_float(row.get("height")) * scale
        draw.rectangle([x, y, x + w, y + h], outline=color, width=3)
    img.save(out, "JPEG", quality=86, optimize=True)
    return out


def cached_patch_preview(patch_id):
    patch_path = PATCH_DIR / f"{patch_id}.png"
    if not patch_path.exists():
        return None
    pdf_cache_key = hashlib.sha1(f"{patch_path.resolve()}|240|82".encode("utf-8")).hexdigest()[:24]
    cached_preview = PDF_IMAGE_CACHE / f"{pdf_cache_key}.jpg"
    if cached_preview.exists():
        return cached_preview
    DOCX_CACHE.mkdir(parents=True, exist_ok=True)
    out = DOCX_CACHE / f"{patch_id}_thumb.jpg"
    if out.exists() and out.stat().st_mtime >= patch_path.stat().st_mtime:
        return out
    with Image.open(patch_path) as img:
        img = img.convert("RGB")
        img.thumbnail((240, 240), Image.Resampling.LANCZOS)
        img.save(out, "JPEG", quality=82, optimize=True)
    return out


def add_label_legend(doc):
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_width(table, [2350, 2350, 2350, 2350])
    for idx, label in enumerate(["oscc", "with_dysplasia", "without_dysplasia", "unknown"]):
        cell = table.cell(0, idx)
        cell.text = display_label(label)
        set_cell_shading(cell, "%02X%02X%02X" % LABEL_COLORS[label])
        style_cell_text(cell, bold=True, size=7.5, color=(255, 255, 255) if label in {"oscc", "without_dysplasia", "unknown"} else (0, 0, 0))
    doc.add_paragraph()


def add_patch_cell(cell, row):
    for p in cell.paragraphs:
        p.paragraph_format.space_after = Pt(0)
    label = row.get("sab_patch_label") or row.get("ndb_ufes_patch_label") or "unknown"
    status = row.get("patch_label_agreement_status") or ""
    heading = cell.paragraphs[0]
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = heading.add_run(row.get("ndb_patch", "patch"))
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(7)
    run.font.color.rgb = RGBColor(0, 0, 0)
    label_p = cell.add_paragraph()
    label_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if status == "disagrees":
        text = f"NDB patch: {display_label(row.get('ndb_ufes_patch_label'))}\nSAB patch: {display_label(row.get('sab_patch_label'))}"
    else:
        text = display_label(label)
    for idx, line in enumerate(text.splitlines()):
        if idx:
            label_p.add_run().add_break()
        r = label_p.add_run(line)
        r.font.name = "Arial"
        r.font.size = Pt(6.2)
    img_path = cached_patch_preview(row.get("ndb_patch"))
    if img_path and img_path.exists():
        pic_p = cell.add_paragraph()
        pic_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pic_p.add_run().add_picture(str(img_path), width=Cm(2.25))


def add_atlas_wsi(doc, numbering, wsi_id, rows, privacy):
    add_heading(doc, wsi_id, 2)
    public_match = sorted({safe_int_text(row.get("previous_public_ndb_wsi_match")) for row in rows if row.get("previous_public_ndb_wsi_match") not in ("", None, "nan")})
    previous_origins = sorted({safe_int_text(row.get("previous_reconstructed_ndb_origin_id")) for row in rows if row.get("previous_reconstructed_ndb_origin_id") not in ("", None, "nan")})
    ndb_wsi_labels = sorted({display_label(row.get("current_origin_label_from_sab_origin_exact_match")) for row in rows if row.get("current_origin_label_from_sab_origin_exact_match") not in ("", None, "nan")})
    sab_label = display_label(rows[0].get("sab_origin_folder_label_normalized"))
    source = "public NDB-UFES + SAB" if public_match else "SAB-only recovered WSI"
    meta = [
        ["Number of patches", str(len(rows))],
        ["WSI source", source],
        ["SAB WSI label", sab_label],
        ["NDB-UFES WSI label", ", ".join(ndb_wsi_labels) if ndb_wsi_labels else "not available"],
        ["Matched NDB-UFES WSI ID", ", ".join(public_match) if public_match else "not available"],
        ["Previous NDB origin ID(s)", ", ".join(previous_origins) if previous_origins else "not available"],
    ]
    if privacy == "lab":
        meta.append(["SAB hash", rows[0].get("sab_origin_image_id", "not available")])
    add_table(doc, meta, [3000, 6400], header=False)

    context = make_context_image(wsi_id, rows)
    if context:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(context), width=Cm(15.8))
        add_caption(doc, numbering.figure_caption(f"{wsi_id} WSI context"))
    add_label_legend(doc)

    sorted_rows = sorted(rows, key=lambda r: r.get("ndb_patch", ""))
    cols = 5
    patch_table = doc.add_table(rows=(len(sorted_rows) + cols - 1) // cols, cols=cols)
    patch_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    patch_table.autofit = False
    set_table_width(patch_table, [1880] * cols)
    for idx, row in enumerate(sorted_rows):
        r = idx // cols
        c = idx % cols
        add_patch_cell(patch_table.cell(r, c), row)
    doc.add_page_break()


def add_atlas(doc, numbering, private_linkage_rows, inventory_rows, privacy, max_wsi=None):
    add_heading(doc, "Validated WSI atlas", 1)
    add_body(doc, "Each atlas entry shows one validated WSI, the recovered patch boxes, and patch thumbnails. Patch thumbnail labels show the patch ID and one patch diagnosis because NDB-UFES and SAB patch labels agree for all current patches.")
    by_wsi = defaultdict(list)
    for row in private_linkage_rows:
        by_wsi[row["final_validated_wsi_id"]].append(row)
    public_wsi = {row["final_validated_wsi_id"] for row in inventory_rows if row.get("wsi_source") == "both"}
    ordered_both = sorted(wsi_id for wsi_id in by_wsi if wsi_id in public_wsi)
    ordered_sab = sorted(wsi_id for wsi_id in by_wsi if wsi_id not in public_wsi)
    if max_wsi:
        ordered = (ordered_both + ordered_sab)[:max_wsi]
        ordered_both = [w for w in ordered if w in public_wsi]
        ordered_sab = [w for w in ordered if w not in public_wsi]

    add_heading(doc, "Atlas: public NDB-UFES + SAB WSI", 2)
    add_body(doc, "These WSI groups have exact SAB patch containment and an exact public NDB-UFES WSI image match.")
    total = len(ordered_both) + len(ordered_sab)
    done = 0
    for wsi_id in ordered_both:
        done += 1
        if done == 1 or done % 25 == 0 or done == total:
            print(f"[{privacy}] atlas {done}/{total}: {wsi_id}", flush=True)
        add_atlas_wsi(doc, numbering, wsi_id, by_wsi[wsi_id], privacy)

    add_heading(doc, "Atlas: SAB-only recovered WSI", 2)
    add_body(doc, "These WSI groups have exact SAB patch containment but no exact public NDB-UFES WSI file match.")
    for wsi_id in ordered_sab:
        done += 1
        if done == 1 or done % 25 == 0 or done == total:
            print(f"[{privacy}] atlas {done}/{total}: {wsi_id}", flush=True)
        add_atlas_wsi(doc, numbering, wsi_id, by_wsi[wsi_id], privacy)


def add_lab_crosswalk(doc, numbering):
    rows = read_csv_rows(CROSSWALK_CSV)
    add_heading(doc, "Laboratory crosswalk", 1)
    add_body(doc, "This lab-only appendix contains original SAB identifiers and must not be included in the public version.")
    table_rows = [["Public WSI pseudonym", "SAB hash name", "SAB folder", "SAB file path"]]
    for row in rows:
        table_rows.append([
            row.get("public_wsi_pseudonym", ""),
            row.get("sab_origin_image_id", ""),
            row.get("sab_origin_folder", ""),
            row.get("sab_origin_path", ""),
        ])
    add_caption(doc, numbering.table_caption("Private WSI crosswalk"))
    add_table(doc, table_rows, [2100, 3200, 1500, 2600])


def add_footer(doc, privacy):
    section = doc.sections[0]
    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run(f"NDB-UFES/SAB validated linkage factsheet - {privacy} version")
    run.font.name = "Arial"
    run.font.size = Pt(7)
    run.font.color.rgb = RGBColor(85, 85, 85)


def build_docx(privacy, max_wsi=None):
    summary = read_json(SUMMARY_JSON)
    linkage_rows = read_csv_rows(LINKAGE_CSV)
    inventory_rows = read_csv_rows(INVENTORY_CSV)
    private_linkage_rows = read_csv_rows(PRIVATE_LINKAGE_CSV)
    doc = Document()
    configure_document(doc)
    add_footer(doc, privacy)
    numbering = Numbering()

    add_cover(doc, summary, privacy)
    add_static_toc(doc, [
        ("Validated linkage summary", 1),
        ("Dataset sources", 1),
        ("Linkage evidence and patch labels", 1),
        ("Validated WSI inventory", 1),
        ("Same-WSI patch similarity", 1),
        ("Validated WSI atlas", 1),
        ("Laboratory crosswalk", 1),
    ] if privacy == "lab" else [
        ("Validated linkage summary", 1),
        ("Dataset sources", 1),
        ("Linkage evidence and patch labels", 1),
        ("Validated WSI inventory", 1),
        ("Same-WSI patch similarity", 1),
        ("Validated WSI atlas", 1),
    ])
    doc.add_page_break()
    add_intro_sections(doc, numbering, summary, linkage_rows, inventory_rows)
    add_similarity_sections(doc, numbering)
    add_atlas(doc, numbering, private_linkage_rows, inventory_rows, privacy, max_wsi=max_wsi)
    if privacy == "lab":
        add_lab_crosswalk(doc, numbering)

    out = VALIDATED_DIR / f"validated_patch_wsi_linkage_factsheet_{privacy}.docx"
    print(f"[{privacy}] saving {out}", flush=True)
    doc.save(out)
    print(f"[{privacy}] saved {out}", flush=True)
    return out


def main():
    parser = argparse.ArgumentParser(description="Build editable DOCX factsheets for validated patch-to-WSI linkage.")
    parser.add_argument("--privacy", choices=["public", "lab", "both"], default="both")
    parser.add_argument("--max-wsi", type=int, default=None, help="Optional limit for quick QA builds.")
    args = parser.parse_args()
    outputs = []
    privacies = ["public", "lab"] if args.privacy == "both" else [args.privacy]
    for privacy in privacies:
        outputs.append(build_docx(privacy, max_wsi=args.max_wsi))
    for output in outputs:
        print(output)


if __name__ == "__main__":
    main()
